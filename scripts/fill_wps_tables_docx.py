"""
Populate "WPS Table Sheels.docx" with the fitted coefficients.

Reads data/generated/paper_table_params.json (written by paper_table_models.py) and
writes each coefficient's b / SE into its model column. Coefficient rows are filled
by model→column position (M1=cols 1-2, M2=cols 3-4, M3=cols 5-6) rather than by
hunting placeholder cells, so terms present in a column the template left blank
(e.g. cubic time2/time3 in inspections M2/M3) still get populated. Rows for terms
no longer in any model (e.g. the ×time interactions) are deleted. Variance rows
(State-to-State σ² and Δσ²) are still filled by placeholder position.

Output: docs/WPS_Table_Sheels_filled_cubic.docx  (the source in ~/Downloads is untouched).
"""
import json
from docx import Document

SRC = '/Users/keshavgoel/Downloads/WPS Table Sheels.docx'
OUT = '/Users/keshavgoel/Research/docs/WPS_Table_Sheels_filled_cubic.docx'
JSON = '/Users/keshavgoel/Research/data/generated/paper_table_params.json'

with open(JSON) as f:
    P = json.load(f)

MODELS = ['M1', 'M2', 'M3']
PLACEHOLDERS = {'X.XX', 'XX.XX', 'XX.X%'}

# label (stripped) -> JSON term key
LABEL_TO_TERM = {
    'Inspections': 'inspections',
    'Time': 'time', 'Time2': 'time2', 'Time3': 'time3',
    'Spending/applicator': 'SPEND_APP_z',
    'Spending/applicator*time': 'SPEND_APP_z:time',
    'Spending/farmworker': 'SPEND_WORK_z',
    'Spending/farmworker*time': 'SPEND_WORK_z:time',
    'Labor Intensity': 'lii_2017_z',
    'H-2A farmworkers:BLS farmworkers': 'h2a_per_farmworker_z',
    'H-2A Authorized/H-2A Requested': 'dol_demand_met_pct_z',
    '%H-2A Authorized to FLC': 'pct_flc_z',
    '%H-2A Authorized to FLC*time': 'pct_flc_z:time',
}


# coefficient columns per model: (b column index, SE column index) within the row
COEF_COLS = {'M1': (1, 2), 'M2': (3, 4), 'M3': (5, 6)}


def term_present(tab, term):
    """True if any model reports this term."""
    return any(tab[m].get(term) is not None for m in MODELS)


def fill_coef_row(row, tab, term):
    """Write b / SE into each model's own columns (position-based, not placeholder-
    based), so cells the template left blank for a present term still get filled."""
    for m in MODELS:
        cell = tab[m].get(term)
        if cell is None:
            continue                       # term absent for this model → leave blank
        bcol, secol = COEF_COLS[m]
        row.cells[bcol].text = f"{cell['b']:.2f}{cell['stars']}"
        row.cells[secol].text = f"{cell['se']:.2f}"


def is_placeholder(text):
    t = text.strip()
    return t in PLACEHOLDERS or t == 'X.XX'


def fill_row(row, values):
    """Fill the row's placeholder cells left-to-right with `values`."""
    targets = [c for c in row.cells if is_placeholder(c.text)]
    if len(targets) != len(values):
        raise SystemExit(
            f"  MISMATCH in row '{row.cells[0].text.strip()}': "
            f"{len(targets)} placeholder cells but {len(values)} values")
    for cell, val in zip(targets, values):
        cell.text = val


def remove_row(row):
    """Delete a table row (used for terms no longer in the model)."""
    row._element.getparent().remove(row._element)


def fill_table(table, tab):
    for row in list(table.rows):
        label = row.cells[0].text.strip()
        if not label:
            continue
        if label.startswith('Δ'):                       # Δ State-to-State σ²
            fill_row(row, [f"{tab[m]['delta_pct']:.1f}%"
                           for m in MODELS if 'delta_pct' in tab[m]])
        elif 'State-to-State' in label:                       # State-to-State σ²
            fill_row(row, [f"{tab[m]['sigma2']:.2f}"
                           for m in MODELS if 'sigma2' in tab[m]])
        elif label in LABEL_TO_TERM:
            term = LABEL_TO_TERM[label]
            if term_present(tab, term):
                fill_coef_row(row, tab, term)
            else:               # term dropped from model (e.g. ×time terms) → drop row
                remove_row(row)


doc = Document(SRC)

# Correct the study period in the titles (data runs 2011–2019, not 2021).
for p in doc.paragraphs:
    if '2011 to 2021' in p.text:
        for run in p.runs:
            run.text = run.text.replace('2011 to 2021', '2011 to 2019')

# doc.tables[0] = inspections (16 rows); doc.tables[1] = violations (20 rows)
fill_table(doc.tables[0], P['inspections'])
fill_table(doc.tables[1], P['violations'])

# Append a collinearity note at the end (full discussion in docs/collinearity_notes.md).
doc.add_paragraph()
note = doc.add_paragraph()
note.add_run('Note on collinearity. ').bold = True
note.add_run(
    'The Labor Intensity Index is entered for the 2017 Census wave only; the '
    '2012, 2017, and 2022 waves are near-collinear (r ≈ 0.977) and produce '
    'unstable, sign-flipping estimates when entered jointly. Spending/applicator '
    'and Spending/farmworker share a common numerator (state STAG obligations) '
    'and are moderately correlated; both are retained as the two FIFRA-protected '
    'populations, but their standard errors should be read with that overlap in '
    'mind. The H-2A block variables are expressed as ratios (H-2A relative to the '
    'BLS farmworker base; authorized relative to requested) to avoid the scale '
    'collinearity of raw counts. See docs/collinearity_notes.md for detail.')

doc.save(OUT)
print(f"Saved filled tables to {OUT}")
